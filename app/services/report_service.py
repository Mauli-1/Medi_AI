from datetime import datetime, timedelta, date
from app.extensions import db
from sqlalchemy import text, func

def daily_sales_report(business_id, branch_id, date):
    sql = text("""
        SELECT s.invoice_number, s.invoice_date, p.full_name as patient,
               s.total_amount, s.payment_mode, s.status
        FROM sales s LEFT JOIN patients p ON s.patient_id = p.id
        WHERE s.business_id = :bid AND DATE(s.invoice_date) = :date
        ORDER BY s.invoice_date DESC
    """)
    rows = db.session.execute(sql, {'bid': business_id, 'date': date}).fetchall()
    return [dict(r._mapping) for r in rows]

def stock_summary(business_id, branch_id=None):
    sql = text("""
        SELECT m.medicine_code, m.name, m.generic_name, m.schedule_type,
               mc.name as category,
               COALESCE(SUM(mb.quantity),0) as total_qty,
               COALESCE(SUM(mb.reserved_qty),0) as reserved_qty,
               m.reorder_level,
               CASE WHEN COALESCE(SUM(mb.quantity),0) <= m.reorder_level THEN 'low'
                    WHEN COALESCE(SUM(mb.quantity),0) >= m.max_stock_level THEN 'over'
                    ELSE 'ok' END as status
        FROM medicines m
        LEFT JOIN medicine_categories mc ON m.category_id = mc.id
        LEFT JOIN medicine_batches mb ON m.id = mb.medicine_id AND mb.is_expired = 0 AND mb.is_damaged = 0
        WHERE m.business_id = :bid AND m.is_deleted = 0
        GROUP BY m.id
        ORDER BY m.name
    """)
    rows = db.session.execute(sql, {'bid': business_id}).fetchall()
    return [dict(r._mapping) for r in rows]

def purchase_register(business_id, branch_id, from_date, to_date):
    sql = text("""
        SELECT p.bill_number, p.bill_date, s.name as supplier,
               p.invoice_number, p.total_amount, p.paid_amount, p.due_amount, p.payment_status
        FROM purchases p JOIN suppliers s ON p.supplier_id = s.id
        WHERE p.business_id = :bid AND p.bill_date BETWEEN :fd AND :td
        ORDER BY p.bill_date DESC
    """)
    rows = db.session.execute(sql, {'bid': business_id, 'fd': from_date, 'td': to_date}).fetchall()
    return [dict(r._mapping) for r in rows]

def gstr1_data(business_id, period):
    sql = text("""
        SELECT s.invoice_number, s.invoice_date, s.customer_gstin,
               s.subtotal, s.cgst_amount, s.sgst_amount, s.igst_amount, s.total_amount
        FROM sales s
        WHERE s.business_id = :bid AND DATE_FORMAT(s.invoice_date,'%Y-%m') = :period
          AND s.status = 'confirmed'
        ORDER BY s.invoice_date
    """)
    rows = db.session.execute(sql, {'bid': business_id, 'period': period}).fetchall()
    return [dict(r._mapping) for r in rows]

def pl_statement(business_id, from_date, to_date):
    rev = db.session.execute(text(
        "SELECT COALESCE(SUM(total_amount),0) as revenue FROM sales WHERE business_id=:bid AND DATE(invoice_date) BETWEEN :fd AND :td AND status='confirmed'"
    ), {'bid': business_id, 'fd': from_date, 'td': to_date}).scalar()
    cogs = db.session.execute(text(
        "SELECT COALESCE(SUM(pi.purchase_rate * si.quantity),0) FROM sale_items si JOIN medicine_batches mb ON si.batch_id=mb.id JOIN purchase_items pi ON mb.purchase_item_id=pi.id JOIN sales s ON si.sale_id=s.id WHERE s.business_id=:bid AND DATE(s.invoice_date) BETWEEN :fd AND :td"
    ), {'bid': business_id, 'fd': from_date, 'td': to_date}).scalar() or 0
    gross = float(rev or 0) - float(cogs)
    return {'revenue': float(rev or 0), 'cogs': float(cogs), 'gross_profit': gross, 'expenses': 0, 'net_profit': gross}

def monthly_sales_report(business_id, year):
    from app.models.sales import Sale
    rows = (
        db.session.query(func.month(Sale.invoice_date).label('month'), func.coalesce(func.sum(Sale.total_amount), 0).label('total'))
        .filter(Sale.business_id == business_id, func.year(Sale.invoice_date) == year, Sale.status == 'confirmed')
        .group_by(func.month(Sale.invoice_date))
        .order_by(func.month(Sale.invoice_date))
        .all()
    )
    return [{'month': r.month, 'total': float(r.total)} for r in rows]


def medicine_wise_sales(business_id, from_date, to_date):
    from app.models.sales import Sale, SaleItem
    from app.models.medicine import Medicine
    rows = (
        db.session.query(Medicine.name, func.sum(SaleItem.quantity).label('qty'), func.sum(SaleItem.total_amount).label('amount'))
        .join(Sale, SaleItem.sale_id == Sale.id)
        .join(Medicine, SaleItem.medicine_id == Medicine.id)
        .filter(Sale.business_id == business_id, func.date(Sale.invoice_date).between(from_date, to_date), Sale.status == 'confirmed')
        .group_by(Medicine.id)
        .order_by(func.sum(SaleItem.total_amount).desc())
        .all()
    )
    return [{'medicine': r.name, 'quantity': int(r.qty), 'amount': float(r.amount)} for r in rows]


def customer_wise_sales(business_id, from_date, to_date):
    from app.models.sales import Sale
    from app.models.patient import Patient
    rows = (
        db.session.query(Patient.full_name, func.count(Sale.id).label('visits'), func.sum(Sale.total_amount).label('total'))
        .join(Sale, Sale.patient_id == Patient.id)
        .filter(Sale.business_id == business_id, func.date(Sale.invoice_date).between(from_date, to_date), Sale.status == 'confirmed')
        .group_by(Patient.id)
        .order_by(func.sum(Sale.total_amount).desc())
        .all()
    )
    return [{'customer': r.full_name, 'visits': r.visits, 'total': float(r.total)} for r in rows]


def cashier_wise_sales(business_id, from_date, to_date):
    from app.models.sales import Sale
    from app.models.user import User
    rows = (
        db.session.query(User.full_name, func.count(Sale.id).label('bills'), func.sum(Sale.total_amount).label('total'))
        .join(Sale, Sale.cashier_id == User.id)
        .filter(Sale.business_id == business_id, func.date(Sale.invoice_date).between(from_date, to_date), Sale.status == 'confirmed')
        .group_by(User.id)
        .order_by(func.sum(Sale.total_amount).desc())
        .all()
    )
    return [{'cashier': r.full_name, 'bills': r.bills, 'total': float(r.total)} for r in rows]


def tax_wise_sales(business_id, from_date, to_date):
    from app.models.sales import Sale, SaleItem
    rows = (
        db.session.query(SaleItem.gst_rate, func.sum(SaleItem.total_amount).label('total'))
        .join(Sale, SaleItem.sale_id == Sale.id)
        .filter(Sale.business_id == business_id, func.date(Sale.invoice_date).between(from_date, to_date), Sale.status == 'confirmed')
        .group_by(SaleItem.gst_rate)
        .order_by(SaleItem.gst_rate)
        .all()
    )
    return [{'gst_rate': float(r.gst_rate), 'total': float(r.total)} for r in rows]


def supplier_wise_purchase(business_id, from_date, to_date):
    from app.models.purchase import Purchase
    from app.models.supplier import Supplier
    rows = (
        db.session.query(Supplier.name, func.count(Purchase.id).label('bills'), func.sum(Purchase.total_amount).label('total'))
        .join(Purchase, Purchase.supplier_id == Supplier.id)
        .filter(Purchase.business_id == business_id, Purchase.bill_date.between(from_date, to_date))
        .group_by(Supplier.id)
        .order_by(func.sum(Purchase.total_amount).desc())
        .all()
    )
    return [{'supplier': r.name, 'bills': r.bills, 'total': float(r.total)} for r in rows]


def product_wise_purchase(business_id, from_date, to_date):
    from app.models.purchase import Purchase, PurchaseItem
    from app.models.medicine import Medicine
    rows = (
        db.session.query(Medicine.name, func.sum(PurchaseItem.quantity).label('qty'), func.sum(PurchaseItem.total_amount).label('amount'))
        .join(Purchase, PurchaseItem.purchase_id == Purchase.id)
        .join(Medicine, PurchaseItem.medicine_id == Medicine.id)
        .filter(Purchase.business_id == business_id, Purchase.bill_date.between(from_date, to_date))
        .group_by(Medicine.id)
        .order_by(func.sum(PurchaseItem.total_amount).desc())
        .all()
    )
    return [{'medicine': r.name, 'quantity': int(r.qty), 'amount': float(r.amount)} for r in rows]


def purchase_audit_report(business_id):
    from app.models.purchase import Purchase
    from app.models.supplier import Supplier
    rows = (
        db.session.query(Purchase.bill_number, Purchase.bill_date, Supplier.name, Purchase.total_amount, Purchase.paid_amount, Purchase.due_amount, Purchase.payment_status)
        .join(Supplier, Purchase.supplier_id == Supplier.id)
        .filter(Purchase.business_id == business_id, Purchase.payment_status != 'paid')
        .order_by(Purchase.bill_date.desc())
        .all()
    )
    return [{'bill_number': r[0], 'bill_date': r[1].isoformat() if r[1] else None, 'supplier': r[2],
              'total': float(r[3]), 'paid': float(r[4]), 'due': float(r[5]), 'status': r[6]} for r in rows]


def expiry_near_report(business_id, days=30):
    from app.models.medicine import Medicine, MedicineBatch
    cutoff = date.today() + timedelta(days=days)
    rows = (
        db.session.query(Medicine.name, MedicineBatch.batch_number, MedicineBatch.expiry_date, MedicineBatch.quantity)
        .join(Medicine, MedicineBatch.medicine_id == Medicine.id)
        .filter(Medicine.business_id == business_id, MedicineBatch.expiry_date <= cutoff, MedicineBatch.expiry_date >= date.today(), MedicineBatch.quantity > 0)
        .order_by(MedicineBatch.expiry_date)
        .all()
    )
    return [{'medicine': r[0], 'batch_number': r[1], 'expiry_date': r[2].isoformat(), 'quantity': r[3]} for r in rows]


def batch_wise_stock(business_id, branch_id=None):
    from app.models.medicine import Medicine, MedicineBatch
    query = (
        db.session.query(Medicine.name, MedicineBatch.batch_number, MedicineBatch.expiry_date, MedicineBatch.quantity, MedicineBatch.rack_location)
        .join(Medicine, MedicineBatch.medicine_id == Medicine.id)
        .filter(Medicine.business_id == business_id, MedicineBatch.quantity > 0)
    )
    if branch_id:
        query = query.filter(MedicineBatch.branch_id == branch_id)
    rows = query.order_by(Medicine.name, MedicineBatch.expiry_date).all()
    return [{'medicine': r[0], 'batch_number': r[1], 'expiry_date': r[2].isoformat() if r[2] else None, 'quantity': r[3], 'rack': r[4]} for r in rows]


def rack_wise_stock(business_id, branch_id=None):
    from app.models.medicine import Medicine, MedicineBatch
    query = (
        db.session.query(MedicineBatch.rack_location, func.sum(MedicineBatch.quantity).label('qty'), func.count(func.distinct(Medicine.id)).label('items'))
        .join(Medicine, MedicineBatch.medicine_id == Medicine.id)
        .filter(Medicine.business_id == business_id, MedicineBatch.quantity > 0)
    )
    if branch_id:
        query = query.filter(MedicineBatch.branch_id == branch_id)
    rows = query.group_by(MedicineBatch.rack_location).order_by(MedicineBatch.rack_location).all()
    return [{'rack': r[0] or 'Unassigned', 'quantity': int(r[1]), 'item_count': r[2]} for r in rows]


def overstock_report(business_id):
    from app.models.medicine import Medicine, MedicineBatch
    rows = (
        db.session.query(Medicine.name, Medicine.max_stock_level, func.coalesce(func.sum(MedicineBatch.quantity), 0).label('qty'))
        .outerjoin(MedicineBatch, db.and_(MedicineBatch.medicine_id == Medicine.id, MedicineBatch.is_expired == False))
        .filter(Medicine.business_id == business_id, Medicine.is_deleted == False)
        .group_by(Medicine.id)
        .having(func.coalesce(func.sum(MedicineBatch.quantity), 0) >= Medicine.max_stock_level)
        .all()
    )
    return [{'medicine': r[0], 'max_stock_level': r[1], 'current_qty': int(r[2])} for r in rows]


def understock_report(business_id):
    from app.models.medicine import Medicine, MedicineBatch
    rows = (
        db.session.query(Medicine.name, Medicine.reorder_level, func.coalesce(func.sum(MedicineBatch.quantity), 0).label('qty'))
        .outerjoin(MedicineBatch, db.and_(MedicineBatch.medicine_id == Medicine.id, MedicineBatch.is_expired == False))
        .filter(Medicine.business_id == business_id, Medicine.is_deleted == False)
        .group_by(Medicine.id)
        .having(func.coalesce(func.sum(MedicineBatch.quantity), 0) <= Medicine.reorder_level)
        .all()
    )
    return [{'medicine': r[0], 'reorder_level': r[1], 'current_qty': int(r[2])} for r in rows]


def schedule_h_register(business_id, from_date, to_date):
    from app.models.compliance import ComplianceRecord
    rows = (
        ComplianceRecord.query.filter(
            ComplianceRecord.business_id == business_id,
            ComplianceRecord.record_type.in_(['schedule_h', 'schedule_h1']),
            ComplianceRecord.dispensed_date >= from_date, ComplianceRecord.dispensed_date <= to_date,
        ).order_by(ComplianceRecord.dispensed_date.desc()).all()
    )
    return rows


def narcotic_register(business_id, from_date, to_date):
    from app.models.compliance import ComplianceRecord
    rows = (
        ComplianceRecord.query.filter(
            ComplianceRecord.business_id == business_id,
            ComplianceRecord.record_type == 'narcotic',
            ComplianceRecord.dispensed_date >= from_date, ComplianceRecord.dispensed_date <= to_date,
        ).order_by(ComplianceRecord.dispensed_date.desc()).all()
    )
    return rows


def patient_purchase_history(patient_id):
    from app.models.sales import Sale
    return Sale.query.filter_by(patient_id=patient_id, status='confirmed').order_by(Sale.invoice_date.desc()).all()


def top_prescribing_doctors(business_id, limit=10):
    from app.models.prescription import Prescription
    from app.models.doctor import Doctor
    rows = (
        db.session.query(Doctor.full_name, func.count(Prescription.id).label('count'))
        .join(Prescription, Prescription.doctor_id == Doctor.id)
        .filter(Doctor.business_id == business_id)
        .group_by(Doctor.id)
        .order_by(func.count(Prescription.id).desc())
        .limit(limit)
        .all()
    )
    return [{'doctor': r[0], 'prescription_count': r[1]} for r in rows]


def export_to_excel(data, headers, sheet_name='Report'):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='1A6B5A')
        cell.font = Font(bold=True, color='FFFFFF')
    for row in data:
        if isinstance(row, dict):
            ws.append([row.get(h, '') for h in headers])
        else:
            ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
